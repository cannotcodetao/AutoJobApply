#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""AutoJobApply · 投递统计表生成器。

数据源只有一个：``state/applications.json``（人工维护的投递台账）。
发送日志 ``state/send_log.json`` 用来做**交叉校验**——同一封邮件记录的工作地点，
在台账里必须一致，否则报错。这条校验是必需的：一旦统计口径和实际发送不一致，
后面所有「投了多少家、都在哪些城市」的判断都会失准。

输出
----
* ``state/applications.xlsx`` —— 带冻结窗格、自动筛选、状态着色的表格
* ``state/applications.md``   —— 纯文本版本，便于 diff 与粘贴

用法
----
::

    python scripts/make_stats.py
    python scripts/make_stats.py --out-xlsx state/我的投递.xlsx

依赖 ``openpyxl``（可选）。缺失时只生成 Markdown，并提示安装命令。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _config

HEADERS = [
    "序号", "公司名称", "业务范围", "投递岗位", "工作地点",
    "投递渠道", "投递日期", "邮箱送达", "当前状态", "备注",
]

# 状态 -> 底色（浅色，xlsx 是白底，不用暗色主题配色）
STATUS_FILLS = {
    "已发送": "E8F5E9",
    "已发送（重投）": "FFF8E1",
    "已收到回执": "E3F2FD",
    "已面试": "E8F5E9",
    "待跟进": "F5F5F5",
    "已退信": "FFEBEE",
}


def load_ledger(cfg) -> dict:
    path = _config.path_of(cfg, "paths.applications", "state/applications.json")
    if not path.exists():
        raise _config.ConfigError(
            f"台账文件不存在：{path}\n"
            "请先创建 state/applications.json，格式见 docs/01-workflow.md「统计台账」一节。"
        )
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise _config.ConfigError(f"台账不是合法 JSON：{path}\n  第 {exc.lineno} 行：{exc.msg}") from exc


def load_send_log(cfg) -> dict:
    path = _config.path_of(cfg, "paths.send_log", "state/send_log.json")
    if not path.exists():
        return {"sends": []}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {"sends": []}


# ---------------------------------------------------------------- 校验


def check_business_length(records, limit: int) -> list:
    """业务范围要有字数上限，否则表格不可读。"""
    problems = []
    for item in records:
        business = str(item.get("business", "") or "")
        if len(business) > limit:
            problems.append(
                f"{item.get('company', '?')}：业务范围 {len(business)} 字，上限 {limit} 字 → 「{business}」"
            )
    return problems


def check_against_log(records, send_log, *, strict_location: bool = True) -> tuple:
    """台账与发送日志交叉校验。

    返回 ``(errors, notices)``：

    * **errors** —— 会导致统计口径失真的问题：工作地点不一致、日志里有台账却没登记。
    * **notices** —— 值得看一眼但不算错：台账登记了 draft_id 但日志里没有对应记录
      （可能是网申渠道、历史台账、或日志被清理过）。
    """
    errors, notices = [], []
    by_id = {record.get("id"): record for record in send_log.get("sends", [])}

    for item in records:
        draft_id = str(item.get("draft_id", "") or "").strip()
        if not draft_id:
            continue

        record = by_id.get(draft_id)
        if not record:
            notices.append(
                f"{draft_id}：台账有登记，但发送日志里没有这封邮件"
                f"（若为网申渠道或日志已清理，可忽略）"
            )
            continue

        log_location = str(record.get("location", "") or "").strip()
        ledger_location = str(item.get("location", "") or "").strip()

        if strict_location and log_location and ledger_location and log_location != ledger_location:
            errors.append(
                f"{draft_id}：工作地点不一致 —— 发送日志「{log_location}」 vs 台账「{ledger_location}」"
            )

    # 反向：日志里有、台账没登记 —— 这是真的漏账，会让统计偏少
    ledger_ids = {str(item.get("draft_id", "")).strip() for item in records if item.get("draft_id")}
    for record in send_log.get("sends", []):
        if record.get("id") and record["id"] not in ledger_ids:
            errors.append(f"{record['id']}：已发送，但台账里没有登记这一条")

    return errors, notices


# ---------------------------------------------------------------- xlsx


def build_xlsx(records, out_path: Path, *, missing_since: dict = None) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[跳过] 未安装 openpyxl，只生成 Markdown。安装：pip install openpyxl")
        return False

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递统计"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="37474F")
    thin = Side(style="thin", color="CFD8DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for index, item in enumerate(records, start=1):
        row = index + 1
        values = [
            index,
            item.get("company", ""),
            item.get("business", ""),
            item.get("role", ""),
            item.get("location", ""),
            item.get("channel", ""),
            item.get("applied_at", ""),
            item.get("delivery", ""),
            item.get("status", ""),
            item.get("note", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in (2, 3, 4, 8, 10),
                horizontal="center" if column in (1, 6, 7) else "left",
            )

        status = str(item.get("status", ""))
        fill_color = STATUS_FILLS.get(status)
        if fill_color:
            sheet.cell(row=row, column=9).fill = PatternFill("solid", fgColor=fill_color)

    widths = [6, 30, 20, 34, 16, 14, 12, 15, 18, 46]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
    sheet.row_dimensions[1].height = 24

    # 第二个工作表：口径说明
    notes_sheet = workbook.create_sheet("口径说明")
    notes = [
        ("统计口径", "只统计「本人确认投递过」的记录；平台自动创建的非本人选择岗位不计入。"),
        ("业务范围", f"统一控制在 {missing_since or 15} 字以内，保证表格可读。"),
        ("邮箱送达", "区分「已发送」与「已送达」：SMTP 返回成功只代表发出，不代表送达。"),
        ("数据来源", "state/applications.json（台账）+ state/send_log.json（发送事实）。"),
        ("交叉校验", "同一封邮件的「工作地点」在台账与日志中必须一致，否则脚本报错。"),
    ]
    notes_sheet.column_dimensions["A"].width = 14
    notes_sheet.column_dimensions["B"].width = 86
    for row_index, (key, value) in enumerate(notes, start=1):
        notes_sheet.cell(row=row_index, column=1, value=key).font = Font(bold=True)
        cell = notes_sheet.cell(row=row_index, column=2, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return True


# ---------------------------------------------------------------- Markdown


def build_md(records, out_path: Path, *, business_limit: int = 15) -> Path:
    lines = [
        "# 投递统计",
        "",
        f"> 共 **{len(records)} 条**投递记录 / **{len({r.get('company', '') for r in records})} 家**公司",
        ">",
        "> 「邮箱送达」与「当前状态」是两列不同的事：前者说邮件有没有送到，后者说流程走到哪。",
        "",
        "| " + " | ".join(HEADERS) + " |",
        "|" + "|".join(["---"] * len(HEADERS)) + "|",
    ]

    for index, item in enumerate(records, start=1):
        note = str(item.get("note", "") or "").replace("\n", " ").replace("|", "\\|")
        values = [
            str(index),
            item.get("company", ""),
            item.get("business", ""),
            item.get("role", ""),
            item.get("location", ""),
            item.get("channel", ""),
            item.get("applied_at", ""),
            str(item.get("delivery", "") or "").replace("\n", " "),
            item.get("status", ""),
            note,
        ]
        lines.append("| " + " | ".join(values) + " |")

    # 按状态汇总
    lines.extend(["", "## 按状态汇总", "", "| 状态 | 条数 |", "|---|---|"])
    tally = {}
    for item in records:
        status = str(item.get("status", "") or "未标注")
        tally[status] = tally.get(status, 0) + 1
    for status, count in sorted(tally.items(), key=lambda pair: -pair[1]):
        lines.append(f"| {status} | {count} |")

    lines.extend([
        "",
        "---",
        "",
        f"*业务范围字数上限 {business_limit} 字；本表由 `scripts/make_stats.py` 自动生成，勿手工编辑。*",
    ])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out_path


# ---------------------------------------------------------------- 入口


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="生成投递统计表")
    parser.add_argument("--out-xlsx", default="", help="xlsx 输出路径（默认 state/applications.xlsx）")
    parser.add_argument("--out-md", default="", help="md 输出路径（默认 state/applications.md）")
    parser.add_argument("--config", default="", help="指定配置文件")
    args = parser.parse_args(argv)

    try:
        cfg = _config.load_config(args.config or None)
        ledger = load_ledger(cfg)
    except _config.ConfigError as exc:
        return _config.fail(exc)

    records = ledger.get("applications", [])
    if not records:
        print("台账里还没有投递记录。")
        return 0

    business_limit = int(
        ledger.get("meta", {}).get("business_scope_max_chars")
        or _config.get(cfg, "stats.business_scope_max_chars", 15)
        or 15
    )

    # ---- 自检 ----
    length_problems = check_business_length(records, business_limit)
    log_errors, log_notices = check_against_log(records, load_send_log(cfg))

    out_xlsx = Path(args.out_xlsx) if args.out_xlsx else _config.path_of(cfg, "paths.stats_xlsx", "state/applications.xlsx")
    out_md = Path(args.out_md) if args.out_md else out_xlsx.with_suffix(".md")

    build_md(records, out_md, business_limit=business_limit)
    built = build_xlsx(records, out_xlsx, missing_since=business_limit)

    print(f"已生成 Markdown：{out_md}")
    if built:
        print(f"已生成 Excel　 ：{out_xlsx}")
    print(f"共 {len(records)} 条记录 / {len({r.get('company', '') for r in records})} 家公司")

    print()
    print(f"业务范围超 {business_limit} 字的记录：", "无 ✅" if not length_problems else "")
    for problem in length_problems:
        print(f"  · {problem}")

    print("发送日志 ↔ 台账交叉校验：", "已同步 ✅" if not log_errors else "")
    for problem in log_errors:
        print(f"  · {problem}")
    for notice in log_notices:
        print(f"  · （提示）{notice}")

    return 0 if not (length_problems or log_errors) else 1


if __name__ == "__main__":
    sys.exit(main())
